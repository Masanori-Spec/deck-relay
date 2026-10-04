#!/usr/bin/env python3
"""Read-only, independently implemented DeckRelay package/receipt verifier.

Usage: python3 tests/oracle.py INPUT.pptx OUTPUT.pptx RECEIPT.json
No production JavaScript or generated application model is imported.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
import hashlib
import io
import json
from pathlib import Path
import posixpath
import re
import sys
from urllib.parse import unquote, urlsplit
import xml.etree.ElementTree as ET
import zipfile

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
RID = f"{{{R}}}id"
SLIDE_REL = R + "/slide"
CLICK = f"{{{A}}}hlinkClick"
HOVER = f"{{{A}}}hlinkHover"
ACTION = "ppaction://hlinksldjump"
MAX_COMPRESSED = 25 * 1024 * 1024
MAX_EXPANDED = 100 * 1024 * 1024
MAX_ENTRIES = 10000
MAX_XML_DEPTH = 128
SCHEMA = "deck-relay-receipt-v1"


class VerificationError(ValueError):
    pass


def require(condition, message):
    if not condition:
        raise VerificationError(message)


def sha(data):
    return hashlib.sha256(data).hexdigest()


def element_children(node):
    return [child for child in node if isinstance(child.tag, str)]


def at_path(root, path):
    require(isinstance(path, list) and len(path) <= MAX_XML_DEPTH,
            "carrierPath must be a bounded array of element-child indexes")
    node = root
    ancestors = []
    for index in path:
        require(type(index) is int and index >= 0, "carrierPath index must be a nonnegative integer")
        children = element_children(node)
        require(index < len(children), "carrierPath does not exist")
        ancestors.append(node)
        node = children[index]
    return node, ancestors


def nodes_with_paths(root):
    stack = [(root, ())]
    while stack:
        node, path = stack.pop()
        yield node, path
        stack.extend((child, path + (i,)) for i, child in reversed(list(enumerate(element_children(node)))))


@dataclass
class XML:
    root: ET.Element
    namespaces: dict
    outside: tuple


def parse_xml(data, label):
    # Null stripping catches UTF-16/32 declarations too. DTDs are outside the profile.
    upper = data.replace(b"\0", b"").upper()
    require(b"<!DOCTYPE" not in upper and b"<!ENTITY" not in upper,
            f"{label}: DTD/entity declarations are unsupported")
    parser = ET.XMLParser(target=ET.TreeBuilder(insert_comments=True, insert_pis=True))
    scopes, stack, pending, outside = {}, [], [], []
    count = 0
    try:
        it = ET.iterparse(io.BytesIO(data), events=("start", "end", "start-ns", "comment", "pi"), parser=parser)
        for event, item in it:
            if event == "start-ns":
                pending.append(item)
            elif event == "start":
                count += 1
                require(count <= 200000, f"{label}: too many XML elements")
                require(len(stack) < MAX_XML_DEPTH, f"{label}: XML is too deep")
                current = dict(stack[-1]) if stack else {}
                current.update(pending)
                pending.clear()
                scopes[id(item)] = current
                stack.append(current)
            elif event == "end":
                stack.pop()
            elif not stack:
                outside.append((event, item.text, item.tail))
        return XML(it.root, scopes, tuple(outside))
    except ET.ParseError as exc:
        raise VerificationError(f"{label}: malformed XML: {exc}") from exc


def fingerprint(node, xml, overrides=None):
    overrides = overrides or {}
    attrs = dict(node.attrib)
    if id(node) in overrides:
        attrs[RID] = overrides[id(node)]
    return (node.tag, sorted(attrs.items()), node.text, node.tail,
            sorted(xml.namespaces.get(id(node), {}).items()),
            tuple(fingerprint(child, xml, overrides) for child in node))


def rels_part(source):
    directory, basename = posixpath.split(source)
    return posixpath.join(directory, "_rels", basename + ".rels")


def resolve_target(source, target):
    require(isinstance(target, str) and target, "empty internal relationship target")
    require("\\" not in target and not any(ord(c) < 32 for c in target), "unsafe relationship target")
    parsed = urlsplit(target)
    require(not (parsed.scheme or parsed.netloc or parsed.query or parsed.fragment),
            "named-slide target must be an internal part without query/fragment")
    require(not re.search(r"%(?:2f|5c)", parsed.path, re.I), "encoded path separator is unsupported")
    path = unquote(parsed.path, encoding="utf-8", errors="strict")
    resolved = posixpath.normpath(path.lstrip("/") if path.startswith("/") else posixpath.join(posixpath.dirname(source), path))
    require(resolved not in ("", ".", "..") and not resolved.startswith("../"),
            "relationship escapes the package")
    return resolved


class Package:
    def __init__(self, path):
        self.path = Path(path)
        require(self.path.stat().st_size <= MAX_COMPRESSED, f"{self.path}: compressed size limit exceeded")
        self.raw = self.path.read_bytes()
        self.digest = sha(self.raw)
        self.entries = {}
        self.xml_cache = {}
        try:
            with zipfile.ZipFile(io.BytesIO(self.raw)) as archive:
                infos = archive.infolist()
                require(len(infos) <= MAX_ENTRIES, "ZIP entry limit exceeded")
                require(sum(i.file_size for i in infos) <= MAX_EXPANDED, "expanded ZIP size limit exceeded")
                for info in infos:
                    name = info.filename
                    require(name not in self.entries, f"duplicate ZIP entry: {name}")
                    parts = name.rstrip("/").split("/")
                    require(name and not name.startswith("/") and "\\" not in name and
                            all(p not in ("", ".", "..") for p in parts) and
                            not any(ord(c) < 32 for c in name), f"unsafe ZIP entry: {name!r}")
                    require(not info.flag_bits & 1, "encrypted ZIP is unsupported")
                    require(info.compress_type in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED),
                            "unsupported ZIP compression")
                    require((info.external_attr >> 16) & 0o170000 != 0o120000, "symlink ZIP entry is unsupported")
                    data = archive.read(info)  # zipfile independently verifies CRC.
                    require(len(data) == info.file_size, f"ZIP size mismatch: {name}")
                    self.entries[name] = data
        except (zipfile.BadZipFile, RuntimeError, NotImplementedError) as exc:
            raise VerificationError(f"invalid ZIP: {exc}") from exc
        lower_names = [name.lower() for name in self.entries]
        require(not any(n.startswith("_xmlsignatures/") or n.endswith("vbaproject.bin") for n in lower_names),
                "signed/macro-enabled packages are outside this profile")
        for name, data in self.entries.items():
            if name.endswith(".xml") or name.endswith(".rels"):
                self.xml(name)
        content = self.entries.get("[Content_Types].xml", b"").lower()
        require(b"macroenabled" not in content and b"vbaproject" not in content,
                "macro-enabled content types are outside this profile")
        self.main = self.main_part()
        self.slides = self.slide_order()
        self.by_part = {part: sid for sid, part in self.slides}
        self.by_id = dict(self.slides)
        require(len(self.slides) <= 300, "slide limit exceeded")

    def xml(self, part):
        require(part in self.entries, f"missing part: {part}")
        if part not in self.xml_cache:
            self.xml_cache[part] = parse_xml(self.entries[part], part)
        return self.xml_cache[part]

    def relationships(self, part):
        if part not in self.entries:
            return {}
        root = self.xml(part).root
        require(root.tag == f"{{{PKG}}}Relationships", f"invalid relationships root: {part}")
        relationships = {}
        for node in element_children(root):
            require(node.tag == f"{{{PKG}}}Relationship", f"unexpected relationship element: {part}")
            rid = node.get("Id")
            require(rid and rid not in relationships, f"missing/duplicate relationship ID: {part}")
            require(node.get("Type") and node.get("Target"), f"incomplete relationship: {part}")
            relationships[rid] = node
        return relationships

    def main_part(self):
        roots = self.relationships("_rels/.rels")
        office = [r for r in roots.values() if r.get("Type") == R + "/officeDocument"]
        require(len(office) == 1 and office[0].get("TargetMode", "Internal") == "Internal",
                "exactly one internal officeDocument relationship is required")
        return resolve_target("", office[0].get("Target"))

    def slide_order(self):
        root = self.xml(self.main).root
        require(root.tag == f"{{{P}}}presentation", "only transitional PresentationML is supported")
        lists = root.findall(f"{{{P}}}sldIdLst")
        require(len(lists) == 1, "exactly one slide ID list is required")
        relationships = self.relationships(rels_part(self.main))
        order, ids, parts = [], set(), set()
        for node in element_children(lists[0]):
            require(node.tag == f"{{{P}}}sldId", "unexpected slide ID list element")
            sid, rid = node.get("id"), node.get(RID)
            require(sid and sid.isdecimal() and sid not in ids, "missing/duplicate slide ID")
            rel = relationships.get(rid)
            require(rel is not None and rel.get("Type") == SLIDE_REL and rel.get("TargetMode", "Internal") == "Internal",
                    "slide order references an invalid internal slide relationship")
            part = resolve_target(self.main, rel.get("Target"))
            require(part not in parts and self.xml(part).root.tag == f"{{{P}}}sld", "duplicate/invalid slide part")
            ids.add(sid)
            parts.add(part)
            order.append((sid, part))
        return order

    def named_target(self, part, node):
        require(node.tag == CLICK and node.get("action") == ACTION, "selected carrier is not an explicit named-slide click")
        rid = node.get(RID)
        rel = self.relationships(rels_part(part)).get(rid)
        require(rel is not None and rel.get("Type") == SLIDE_REL and rel.get("TargetMode", "Internal") == "Internal",
                "selected carrier does not reference an internal slide")
        target = resolve_target(part, rel.get("Target"))
        require(target in self.by_part, "selected target is absent from the presentation slide order")
        return target


def read_json(path):
    def pairs(items):
        result = {}
        for key, value in items:
            require(key not in result, f"duplicate JSON key: {key}")
            result[key] = value
        return result
    def constant(value):
        raise VerificationError(f"invalid JSON constant: {value}")
    require(Path(path).stat().st_size <= 4 * 1024 * 1024, "receipt is too large")
    return json.loads(Path(path).read_text(encoding="utf-8"), object_pairs_hook=pairs, parse_constant=constant)


def load_receipt(path):
    result = read_json(path)
    require(isinstance(result, dict) and result.get("schema") == SCHEMA, "unsupported receipt schema")
    require(isinstance(result.get("changes"), list) and len(result["changes"]) <= 3000, "invalid changes list")
    require(isinstance(result.get("changedParts"), list), "invalid changedParts list")
    return result


def carrier_kind(node, ancestors):
    require(ancestors, "click carrier cannot be the root")
    parent = ancestors[-1]
    require(sum(c.tag == CLICK for c in element_children(parent)) == 1,
            "duplicate click elements on one carrier are unsupported")
    if parent.tag == f"{{{P}}}cNvPr":
        # p:spTree's own group properties are not a clickable shape.
        require(len(ancestors) >= 3 and ancestors[-3].tag in
                {f"{{{P}}}{name}" for name in ("sp", "pic", "cxnSp", "graphicFrame", "grpSp")},
                "whole-shape click does not belong to a supported shape")
        return "shape"
    require(parent.tag == f"{{{A}}}rPr" and len(ancestors) >= 2 and
            ancestors[-2].tag == f"{{{A}}}r", "click is not on a supported shape/text run")
    return "run"


def python_pptx_check(package):
    try:
        from pptx import Presentation
        from pptx.enum.action import PP_ACTION
    except ImportError as exc:
        raise VerificationError("python-pptx is required; install the pinned verification dependency") from exc
    presentation = Presentation(io.BytesIO(package.raw))
    # Accessing Presentation.slides renames slide part names in memory in
    # python-pptx 1.0.2. Stable slide IDs remain valid; never save this object.
    api_order = [str(s.slide_id) for s in presentation.slides]
    require(api_order == [sid for sid, _ in package.slides], "python-pptx disagrees with independently resolved slide order")
    checked = 0
    expected = 0
    for sid, part in package.slides:
        root = package.xml(part).root
        expected_ids = {}
        for node, path in nodes_with_paths(root):
            if node.tag != CLICK or node.get("action") != ACTION:
                continue
            _, ancestors = at_path(root, list(path))
            if not ancestors or ancestors[-1].tag != f"{{{P}}}cNvPr":
                continue
            carrier_kind(node, ancestors)
            shape_id = ancestors[-1].get("id")
            require(shape_id and shape_id not in expected_ids, "ambiguous whole-shape ID")
            expected_ids[shape_id] = package.named_target(part, node)
        expected += len(expected_ids)
        slide = presentation.slides[api_order.index(sid)]
        seen = set()
        def walk(shapes):
            nonlocal checked
            for shape in shapes:
                shape_id = str(shape.shape_id)
                if shape_id in expected_ids:
                    require(shape_id not in seen, "python-pptx found an ambiguous shape ID")
                    seen.add(shape_id)
                    action = shape.click_action
                    require(action.action == PP_ACTION.NAMED_SLIDE, "python-pptx did not identify a named-slide action")
                    target = action.target_slide
                    require(target is not None and str(target.slide_id) == package.by_part[expected_ids[shape_id]],
                            "python-pptx target disagrees with independent XML resolution")
                    checked += 1
                if hasattr(shape, "shapes"):
                    walk(shape.shapes)
        walk(slide.shapes)
        require(seen == set(expected_ids), "python-pptx did not expose all supported whole-shape carriers")
    require(checked == expected, "whole-shape API coverage is incomplete")
    return checked


def verify(input_path, output_path, receipt_path, expected_path=None):
    before, after = Package(input_path), Package(output_path)
    receipt = load_receipt(receipt_path)
    if expected_path is not None:
        expected = read_json(expected_path)
        fields = {"sourcePart", "sourceSlideId", "carrierPath", "kind", "fromSlideId", "toSlideId"}
        require(isinstance(expected, list) and all(isinstance(e, dict) and fields <= set(e) for e in expected),
                "invalid independent expected-selection fixture")
        def expectation_key(item):
            return json.dumps({field: item.get(field) for field in sorted(fields)}, sort_keys=True, ensure_ascii=False)
        actual_keys = [expectation_key(item) for item in receipt["changes"] if isinstance(item, dict)]
        expected_keys = [expectation_key(item) for item in expected]
        require(len(set(expected_keys)) == len(expected_keys), "duplicate independent expected selection")
        require(sorted(actual_keys) == sorted(expected_keys), "receipt does not match independently authored intended selections")
    require(receipt.get("inputSha256") == before.digest, "input SHA-256 mismatch")
    require(receipt.get("outputSha256") == after.digest, "output SHA-256 mismatch")
    require(receipt["changes"] or before.raw == after.raw, "a wholly no-op result must preserve complete PPTX bytes")
    require(set(before.entries) == set(after.entries), "ZIP entry set changed")
    require(before.main == after.main and before.slides == after.slides, "slide identities/order changed")
    changed = {name for name in before.entries if before.entries[name] != after.entries[name]}
    allowed = set()
    overrides = {}
    selected = set()
    selected_new_rids = {}
    expected_fields = {"sourcePart", "sourceSlideId", "carrierPath", "kind", "oldRid", "newRid",
                       "fromSlideId", "toSlideId", "fromPart", "toPart"}
    for change in receipt["changes"]:
        require(isinstance(change, dict) and expected_fields <= set(change), "incomplete change record")
        part = change["sourcePart"]
        require(isinstance(part, str) and part in before.by_part, "change source is not a presentation slide")
        require(change["sourceSlideId"] == before.by_part[part], "source slide ID/part mismatch")
        old_xml, new_xml = before.xml(part), after.xml(part)
        old, old_ancestors = at_path(old_xml.root, change["carrierPath"])
        new, new_ancestors = at_path(new_xml.root, change["carrierPath"])
        key = (part, tuple(change["carrierPath"]))
        require(key not in selected, "duplicate selected carrier")
        selected.add(key)
        kind = carrier_kind(old, old_ancestors)
        require(change["kind"] == kind == carrier_kind(new, new_ancestors), "carrier kind mismatch")
        old_target, new_target = before.named_target(part, old), after.named_target(part, new)
        require(old_target != new_target, "changes must omit same-target no-ops")
        require(change["fromPart"] == old_target and change["toPart"] == new_target,
                "receipt destination part mismatch")
        require(change["fromSlideId"] == before.by_part[old_target] and
                change["toSlideId"] == after.by_part[new_target], "receipt destination slide ID mismatch")
        require(change["oldRid"] == old.get(RID) and change["newRid"] == new.get(RID), "receipt relationship ID mismatch")
        require(old.get(RID) != new.get(RID), "retargeting must isolate the selected carrier with a different relationship ID")
        overrides.setdefault(part, {})[id(new)] = old.get(RID)
        selected_new_rids.setdefault(part, set()).add(new.get(RID))
        allowed.update((part, rels_part(part)))
    require(changed <= allowed, "an unselected/unrelated part changed: " + ", ".join(sorted(changed - allowed)))
    for part, replacements in overrides.items():
        require(before.xml(part).outside == after.xml(part).outside, f"{part}: out-of-root XML comments/PIs changed")
        require(fingerprint(before.xml(part).root, before.xml(part)) ==
                fingerprint(after.xml(part).root, after.xml(part), replacements),
                f"{part}: XML changed beyond selected click relationship IDs")
    added_relationships = 0
    for part, new_rids in selected_new_rids.items():
        relationship_part = rels_part(part)
        bx, ax = before.xml(relationship_part), after.xml(relationship_part)
        br, ar = before.relationships(relationship_part), after.relationships(relationship_part)
        require(set(br) <= set(ar), "existing relationships were removed")
        added = set(ar) - set(br)
        require(added <= new_rids, "an added relationship is not used by a selected carrier")
        added_relationships += len(added)
        # Remove only verified new relationships from the comparison tree.
        require(bx.outside == ax.outside and bx.root.tag == ax.root.tag and bx.root.attrib == ax.root.attrib and
                bx.root.text == ax.root.text and bx.root.tail == ax.root.tail and
                bx.namespaces[id(bx.root)] == ax.namespaces[id(ax.root)],
                "relationship root changed")
        existing_after = [n for n in ax.root if not (isinstance(n.tag, str) and n.get("Id") in added)]
        require(len(list(bx.root)) == len(existing_after), "relationship children changed")
        require(all(fingerprint(b, bx) == fingerprint(a, ax) for b, a in zip(bx.root, existing_after)),
                "an existing relationship changed (shared-rId isolation failed)")
        # Additions must be append-only, simple internal slide relationships.
        element_ids = [n.get("Id") for n in element_children(ax.root)]
        require(element_ids[:len(br)] == list(br), "existing relationships reordered")
        for rid in added:
            rel = ar[rid]
            require(set(rel.attrib) in ({"Id", "Type", "Target"}, {"Id", "Type", "Target", "TargetMode"}) and
                    rel.get("Type") == SLIDE_REL and rel.get("TargetMode", "Internal") == "Internal" and
                    not list(rel) and not (rel.text or "").strip() and not (rel.tail or "").strip(),
                    "added relationship has unsupported content")
            target = resolve_target(part, rel.get("Target"))
            require(target in after.by_part, "added relationship has a missing target")
    reported = {}
    for item in receipt["changedParts"]:
        require(isinstance(item, dict) and {"part", "beforeSha256", "afterSha256"} <= set(item), "incomplete changed part record")
        name = item["part"]
        require(isinstance(name, str) and name not in reported and name in changed, "duplicate/unexpected changed part record")
        require(item["beforeSha256"] == sha(before.entries[name]) and item["afterSha256"] == sha(after.entries[name]),
                f"part hash mismatch: {name}")
        reported[name] = item
    require(set(reported) == changed, "changedParts does not exactly enumerate actual changed entries")
    # Preservation of all other bytes also covers media, notes, layouts, masters,
    # content types, original relationships, thumbnails and presentation metadata.
    media = {name: sha(data) for name, data in before.entries.items() if "/media/" in name and not name.endswith("/")}
    require(all(sha(after.entries[name]) == digest for name, digest in media.items()), "media bytes changed")
    unselected = 0
    for _, part in before.slides:
        for node, path in nodes_with_paths(before.xml(part).root):
            if node.tag in (CLICK, HOVER) and (part, path) not in selected:
                unselected += 1
    api_before = python_pptx_check(before)
    api_after = python_pptx_check(after)
    return {"ok": True, "schema": SCHEMA, "intendedSelectionVerified": expected_path is not None, "slides": len(before.slides),
            "selectedChanges": len(selected), "unselectedClickAndHoverPreserved": unselected,
            "changedParts": sorted(changed), "unchangedEntries": len(before.entries) - len(changed),
            "mediaSha256": media, "addedRelationships": added_relationships,
            "pythonPptxWholeShapeChecks": {"input": api_before, "output": api_after},
            "nativePowerPointTested": False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("receipt", type=Path)
    parser.add_argument("--expected-selection", type=Path, help="independently authored expected selections, recommended for fixtures")
    args = parser.parse_args()
    try:
        result = verify(args.input, args.output, args.receipt, args.expected_selection)
    except Exception as exc:
        # Any unsupported package/API condition is a failure, never a skipped pass.
        print(json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False), file=sys.stderr)
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
